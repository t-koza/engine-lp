"""出力: CSV / Excel(xlsx) / 取得ログ / 取得失敗一覧。"""
from __future__ import annotations

import csv
import datetime as _dt
from typing import Dict, List

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from extract import COLUMNS, StoreRecord


def write_csv(path: str, records: List[StoreRecord]) -> None:
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(COLUMNS)
        for r in records:
            w.writerow(r.as_row())


def write_xlsx(path: str, records: List[StoreRecord]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "店舗リスト"
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    ws.append(COLUMNS)
    for c in range(1, len(COLUMNS) + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for r in records:
        ws.append(r.as_row())
    widths = [18, 22, 10, 36, 40, 40, 16, 16, 26, 26, 30, 30, 14, 12, 30]
    for i, wdt in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = wdt
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{ws.max_row}"
    wb.save(path)


def write_fetch_log(path: str, rows: List[Dict]) -> None:
    cols = ["timestamp", "company", "url", "kind", "status", "ok", "redirected",
            "cross_host_redirect", "final_url", "bytes", "elapsed_ms", "note"]
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=cols, extrasaction="ignore")
        w.writeheader()
        for row in rows:
            w.writerow(row)


def write_failed(path: str, rows: List[Dict]) -> None:
    cols = ["company", "prefecture", "url", "reason"]
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=cols, extrasaction="ignore")
        w.writeheader()
        for row in rows:
            w.writerow(row)


def write_summary(path: str, stats: Dict) -> None:
    now = _dt.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    lines = [
        "# 取得サマリ",
        "",
        f"- 実行日時: {now}",
        f"- 出力店舗件数(重複除去後): **{stats.get('records', 0)}**",
        f"- HTTP200確認済み 店舗URL: {stats.get('verified_store_urls', 0)}",
        f"- HTTP200確認済み 問い合わせフォームURL: {stats.get('verified_forms', 0)}",
        f"- 巡回ページ総数: {stats.get('pages_fetched', 0)}",
        f"- 取得失敗(会社/ロケーター): {stats.get('failed', 0)}",
        "",
        "## 会社別 店舗件数",
        "",
        "| 会社 | 店舗件数 |",
        "| --- | ---: |",
    ]
    for company, n in sorted(stats.get("per_company", {}).items(), key=lambda x: -x[1]):
        lines.append(f"| {company} | {n} |")
    lines += [
        "",
        "## 品質ルール(本実行で適用)",
        "- 店舗名 と (住所 または 店舗URL) の両方が揃わない行は出力しない。",
        "- 全URLは実際にGETし、最終ステータス200のもののみ『確認済み』として備考に明記。",
        "- 404/リンク切れ/別ホストへのリダイレクトは備考に記録し、該当URLは空欄化または要確認表示。",
        "- 重複は (店舗URL正規化) → (会社+住所) → (会社+店舗名+都道府県) の順に判定し1件に統合。",
        "- 取得できない項目は推測せず空欄。",
    ]
    with open(path, "w", encoding="utf-8") as f:
        f.write("\n".join(lines) + "\n")
